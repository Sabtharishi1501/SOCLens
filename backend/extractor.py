"""
Extracts plain text from an uploaded SOP document (.pdf, .docx, or .txt),
cleans it up, and enforces the fixed MAX_INPUT_CHARS limit from config.py.

Design choice: truncation is never silent. If the document is longer than
the limit, we cut it and return truncated=True so the caller can surface
that to the user instead of quietly analyzing a partial document.
"""

import io
import re

import pdfplumber
from docx import Document

from config import MAX_INPUT_CHARS


class ExtractionError(Exception):
    """Raised when a file can't be read or contains no usable text."""


def extract_text(filename: str, file_bytes: bytes) -> dict:
    """
    Returns a dict:
        {
            "text": str,          # cleaned, possibly truncated text
            "truncated": bool,    # True if original text exceeded MAX_INPUT_CHARS
            "original_chars": int,  # length of text before truncation
            "used_chars": int,       # length of text actually sent onward
        }
    Raises ExtractionError on unsupported/unreadable files or empty content.
    """
    ext = _get_extension(filename)

    if ext == "pdf":
        raw_text = _extract_pdf(file_bytes)
    elif ext == "docx":
        raw_text = _extract_docx(file_bytes)
    elif ext == "txt":
        raw_text = _extract_txt(file_bytes)
    else:
        raise ExtractionError(
            f"Unsupported file type '.{ext}'. Please upload a .pdf, .docx, or .txt file."
        )

    cleaned = _clean_text(raw_text)

    if not cleaned:
        raise ExtractionError(
            "No readable text was found in this file. If it's a scanned/image-based "
            "PDF, text extraction won't work without OCR, which this tool doesn't do."
        )

    original_chars = len(cleaned)
    truncated = original_chars > MAX_INPUT_CHARS
    used_text = cleaned[:MAX_INPUT_CHARS] if truncated else cleaned

    return {
        "text": used_text,
        "truncated": truncated,
        "original_chars": original_chars,
        "used_chars": len(used_text),
    }


def _get_extension(filename: str) -> str:
    if "." not in filename:
        return ""
    return filename.rsplit(".", 1)[-1].lower()


def _extract_pdf(file_bytes: bytes) -> str:
    try:
        text_parts = []
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text() or ""
                text_parts.append(page_text)
        return "\n".join(text_parts)
    except Exception as e:
        raise ExtractionError(f"Could not read PDF file: {e}")


def _extract_docx(file_bytes: bytes) -> str:
    try:
        doc = Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs]

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text)

        return "\n".join(paragraphs)
    except Exception as e:
        raise ExtractionError(f"Could not read DOCX file: {e}")


def _extract_txt(file_bytes: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return file_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ExtractionError("Could not decode text file (unrecognized encoding).")


def _clean_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()